#!/usr/bin/env python3
"""Build the Qumaa recruiter guide from screenshots of the actual application.

Run from the repository root:
    python3 presentations/source/build_guide.py

Dependencies: python-pptx, Pillow. No network access is used. Screenshot inputs
must be captured from an isolated, fictional demo library. --allow-missing is
for layout development only; the normal build rejects missing screenshots.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "presentations"
SHOTS = BASE / "screenshots"
OUTPUT = BASE / "Qumaa_Recruiter_Guide.pptx"
VERSION = "v0.1.0-beta.2"
LINKEDIN = "https://www.linkedin.com/in/fawazsalehq/"
# Arial is installed on the presentation host and travels well across viewers.
# Application screenshots retain the product's bundled IBM Plex typography.
FONT = "Arial"
MONO = "Arial"
C = {
    "ink": "2B2733", "soft": "3D3949", "violet": "6D3FB0",
    "lavender": "C0A3EA", "paper": "F4F1F8", "white": "FFFFFF",
    "muted": "716B7B", "line": "DDD5E7", "pale": "ECE4F7",
    "blue": "0A66C2",
}
REQUIRED = [
    "01-jobs.png", "02-library.png", "03-job-editor.png", "04-criteria.png",
    "05-filters.png", "06-evidence.png", "07-review.png", "08-profile.png",
    "09-comparison-header.png", "09-comparison-criteria.png", "10-shortlist.png",
    "11-versions.png", "12-refresh.png", "12-refresh-action.png", "13-express-form.png", "13-express-results.png",
    "14-watched-folder.png", "14-backups.png", "16-status.png", "17-hiring.png",
    "18-ai-settings.png", "19-ai-review.png",
]
# Capture scripts supply focused, ready-to-fit images in each language.
CROPS = {}
ALLOW_MISSING = False


def rgb(value):
    return RGBColor.from_string(C.get(value, value))


def rect(slide, x, y, w, h, fill="white", line=None, rounded=False):
    shape = slide.shapes.add_shape(
        SHAPE.ROUNDED_RECTANGLE if rounded else SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.7)
    else:
        shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = 0.12
    return shape


def text(slide, value, x, y, w, h, size=18, color="ink", bold=False,
         font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0,
         link=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for index, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(3)
        p.line_spacing = 1.1
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        if link:
            run.hyperlink.address = link
    return shape


def number_badge(slide, value, x, y, diameter=.34, size=11):
    """Create a numbered circle with true horizontal and vertical centering."""
    badge = slide.shapes.add_shape(
        SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = rgb("violet")
    badge.line.fill.background()

    # Keep the number in the circle's own text frame. The former separate,
    # shorter textbox used top anchoring and paragraph-after spacing, which
    # shifted the glyph away from the visual center of the badge.
    tf = badge.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.space_before = p.space_after = Pt(0)
    p.line_spacing = 1.0
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = rgb("white")
    return badge


def rule(slide, x, y, w, color="line", width=0.6):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def arrow_line(slide, points, color="violet", width=1.8, dashed=False):
    """Native editable line segments; only the terminal segment has an arrow."""
    for i, (a, b) in enumerate(zip(points, points[1:])):
        shape = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(a[0]), Inches(a[1]), Inches(b[0]), Inches(b[1])
        )
        shape.line.color.rgb = rgb(color)
        shape.line.width = Pt(width)
        ln = shape._element.spPr.get_or_add_ln()
        if dashed:
            dash = OxmlElement("a:prstDash")
            dash.set("val", "dash")
            ln.append(dash)
        if i == len(points) - 2:
            end = OxmlElement("a:tailEnd")
            end.set("type", "triangle")
            end.set("w", "med")
            end.set("len", "med")
            ln.append(end)


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def brand(slide, x=.58, y=.16, dark=False, size=.30):
    icon = ROOT / "build" / "icon.png"
    if icon.exists():
        slide.shapes.add_picture(str(icon), Inches(x), Inches(y), Inches(size), Inches(size))
    text(slide, "QUMAA", x + size + .12, y + .035, 1.7, .21,
         size=10, font=MONO, color="lavender" if dark else "violet", bold=True)


def footer(slide, number, dark=False, label="Actual application · Fictional demo data"):
    tone = "lavender" if dark else "muted"
    rule(slide, .58, 7.04, 12.17, "soft" if dark else "line")
    text(slide, label, .60, 7.13, 8.4, .20, size=9.2, color=tone)
    text(slide, f"Qumaa {VERSION}   /   {number:02d}", 9.1, 7.13, 3.63, .20,
         size=9.2, font=MONO, color=tone, align=PP_ALIGN.RIGHT)


def slide_base(prs, title, number, subtitle=None, dark=False,
               footer_label="Actual application · Fictional demo data"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("ink" if dark else "paper")
    brand(slide, dark=dark)
    text(slide, title, .58, .61, 12.15, .66, size=30, bold=True,
         color="white" if dark else "ink")
    if subtitle:
        text(slide, subtitle, .60, 1.24, 12.0, .36, size=14.5,
             color="lavender" if dark else "muted")
    footer(slide, number, dark, footer_label)
    return slide


def shot(slide, name, x, y, w, h, caption=None, framed=True, crop=None):
    path = SHOTS / name
    if framed:
        rect(slide, x, y, w, h, "white", "line", rounded=False)
    pad = .07 if framed else 0
    left, top, width, height = x + pad, y + pad, w - 2 * pad, h - 2 * pad
    if caption:
        caption_shape = text(slide, caption, x + .05, y - .27, w - .1, .23,
                             size=10.2, color="muted", font=MONO)
        caption_shape.name = "screenshot-caption:" + name
    if not path.exists():
        if not ALLOW_MISSING:
            raise FileNotFoundError(f"Missing actual application screenshot: {path}")
        text(slide, f"SCREENSHOT PENDING\n{name}", left + .1, top + height / 2 - .4,
             width - .2, .8, size=16, color="muted", align=PP_ALIGN.CENTER)
        return
    with Image.open(path) as im:
        iw, ih = im.size
    region = crop or CROPS.get(name) or (0, 0, iw, ih)
    x0, y0, x1, y1 = region
    if not (0 <= x0 < x1 <= iw and 0 <= y0 < y1 <= ih):
        raise ValueError(f"Invalid screenshot crop for {name}: {region} in {(iw, ih)}")
    scale = min(width / (x1 - x0), height / (y1 - y0))
    dw, dh = (x1 - x0) * scale, (y1 - y0) * scale
    picture = slide.shapes.add_picture(str(path), Inches(left + (width - dw) / 2),
                                       Inches(top + (height - dh) / 2), Inches(dw), Inches(dh))
    picture.name = "screenshot:" + name
    picture.crop_left, picture.crop_top = x0 / iw, y0 / ih
    picture.crop_right, picture.crop_bottom = 1 - x1 / iw, 1 - y1 / ih


def steps(slide, items, x=.63, y=1.93, w=3.08, gap=1.35, size=19):
    for i, (heading, body) in enumerate(items):
        yy = y + i * gap
        number_badge(slide, str(i + 1), x, yy + .025)
        text(slide, heading, x + .47, yy, w - .47, .64, size=size, bold=True)
        text(slide, body, x + .47, yy + .64, w - .47, gap - .60, size=15, color="soft")


def callout(slide, value, x=.64, y=6.12, w=3.08, h=.64, size=13.5):
    rect(slide, x, y, w, h, "pale", rounded=True)
    text(slide, value, x + .14, y + .10, w - .28, h - .14, size=size, color="violet")


def standard(prs, number, title, subtitle, screenshot, items, note,
             hint=None, screenshot_caption=None):
    slide = slide_base(prs, title, number, subtitle)
    steps(slide, items)
    if screenshot == "12-refresh.png":
        rect(slide, 4.04, 1.84, 8.71, 4.96, "white", "line", True)
        shot(slide, screenshot, 4.19, 2.50, 8.41, 1.08, framed=False)
        arrow_line(slide, [(8.40, 3.82), (8.40, 4.17)], width=1.6)
        shot(slide, "12-refresh-action.png", 4.19, 4.36, 8.41, 1.58, framed=False)
    else:
        shot(slide, screenshot, 4.04, 1.84, 8.71, 4.96, screenshot_caption)
    if hint:
        callout(slide, hint)
    notes(slide, note)
    return slide


def side_panel(prs, number, title, subtitle, screenshot, items, note, hint=None):
    slide = slide_base(prs, title, number, subtitle)
    steps(slide, items, w=6.1, gap=1.36, size=21)
    shot(slide, screenshot, 7.1, 1.78, 5.65, 5.02)
    if hint:
        callout(slide, hint, w=6.1, size=14)
    notes(slide, note)
    return slide


def make_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = "Qumaa | Recruiter Guide"
    prs.core_properties.subject = "Application use, candidate assessment, and recruiting records"
    prs.core_properties.author = "Fawaz Aljathnan"
    prs.core_properties.keywords = "Qumaa, recruiting, Express search, job status, hiring records, AI reranking"
    prs.core_properties.comments = "Actual application screenshots with fictional demonstration data."

    # 01: introduction, followed immediately by the contents.
    s = slide_base(prs, "Qumaa: application and recruiter guide", 1,
                   "Purpose, workflows, and operating guidance", dark=True)
    text(s, "Application purpose", .68, 1.98, 5.2, .48, size=25, bold=True, color="white")
    text(s, "Search a local CV library against job requirements, review supporting evidence, and maintain recruiting decisions.",
         .68, 2.67, 5.13, 1.80, size=23, color="white")
    text(s, "Demonstration scenario", .68, 4.57, 5.2, .36, size=18, bold=True, color="lavender")
    text(s, "Training & Development Manager\nRiyadh food and beverage group\nFictional candidates and screening records", .68, 5.10, 5.15, 1.23, size=17, color="white")
    shot(s, "01-jobs.png", 6.18, 1.96, 6.54, 4.67)
    notes(s, "Qumaa is a desktop application for recruiters working with an existing CV library. This guide explains how to import CVs, search a saved or temporary job, assess candidates, and record hiring outcomes. Standard matching and recruiting records are local. Optional AI reranking sends a reviewed request to the selected provider. The example vacancy and all candidate information are fictional. Recruiters verify evidence and make the hiring decision.")

    s = slide_base(prs, "Table of contents", 2, "Sections follow the sequence of a recruiting assignment.", footer_label="Recruiter guide")
    toc = [
        ("Workspace and CV library", "Navigation, workflow, and importing CVs", "03–05", 3),
        ("Jobs and search", "Saved jobs, Express search, criteria, and filters", "06–09", 6),
        ("Candidate assessment", "Evidence, reviews, profiles, comparison, and shortlists", "10–15", 10),
        ("Job records", "Refreshing matches, status history, and hiring outcomes", "16–18", 16),
        ("Optional AI reranking", "Providers, data review, and result interpretation", "19–21", 19),
        ("Operations and reference", "Watched folders, backups, and record scope", "22–24", 22),
    ]
    toc_links = []
    for i, (heading, body, pages, target) in enumerate(toc):
        x, y = .65 + (i // 3) * 6.22, 1.93 + (i % 3) * 1.62
        box = rect(s, x, y, 5.95, 1.38, "white", "line", True)
        text(s, heading, x + .20, y + .17, 4.78, .40, size=21, bold=True)
        text(s, pages, x + 4.90, y + .20, .88, .29, size=13.5, color="violet", align=PP_ALIGN.RIGHT)
        text(s, body, x + .20, y + .74, 5.48, .61, size=16, color="soft")
        toc_links.append((box, target))
    notes(s, "Use the page ranges to locate a topic. In the editable presentation, each section card links to its first page. Pages 3 to 5 introduce the workspace and library; 6 to 9 cover job searches; 10 to 15 cover assessment; 16 to 18 cover job records; 19 to 21 explain optional AI; and 22 to 24 cover operations and reference information.")

    s = slide_base(prs, "Workspace navigation", 3, "Jobs contains both saved vacancies and Express search.")
    for i, (heading, body) in enumerate([
        ("Jobs", "Search vacancies; retain reviews, shortlists, and hiring records."),
        ("Library", "Import CVs, inspect candidate details, and link CV versions."),
        ("Settings", "Set preferences, AI providers, watched folders, and backups."),
    ]):
        y = 1.96 + i * 1.54
        rect(s, .65, y, 3.16, 1.44, "white", "line", True)
        text(s, heading, .84, y + .14, 2.77, .34, size=21, bold=True, color="violet")
        text(s, body, .84, y + .53, 2.74, .87, size=15, color="soft")
    shot(s, "01-jobs.png", 4.05, 1.94, 8.67, 4.84)
    notes(s, "The application has three main tabs: Jobs, Library, and Settings. Jobs includes Express search next to New job. Both saved and temporary searches use the same local matching workflow and optional AI reranking. Candidate profiles apply across jobs; vacancy criteria, reviews, shortlists, status history, and hiring outcomes belong to a saved job.")

    s = slide_base(prs, "Recruiting workflow", 4, "Example: screening a Training & Development Manager vacancy.", footer_label="Recruiter workflow · Editable diagram")
    coords = [(.65, 1.97), (4.71, 1.97), (8.77, 1.97), (8.77, 4.05), (4.71, 4.05), (.65, 4.05)]
    cards = [
        ("Import CVs", "Build the local candidate library."),
        ("Define requirements", "Create a saved job or an Express search."),
        ("Find matches", "Review results from the local library."),
        ("Assess candidates", "Check evidence, compare, and record reviews."),
        ("Prepare the shortlist", "Choose candidates, set order, and export."),
        ("Record the outcome", "Update status, hired candidate, and backups."),
    ]
    for i, ((x, y), (heading, body)) in enumerate(zip(coords, cards)):
        rect(s, x, y, 3.80, 1.58, "white", "line", True)
        number_badge(s, str(i + 1), x + .19, y + .20)
        text(s, heading, x + .70, y + .16, 2.86, .57, size=20, bold=True)
        text(s, body, x + .20, y + .85, 3.40, .64, size=16, color="soft")
    arrow_line(s, [(4.46, 2.72), (4.68, 2.72)])
    arrow_line(s, [(8.52, 2.72), (8.74, 2.72)])
    arrow_line(s, [(10.67, 3.60), (10.67, 4.02)])
    arrow_line(s, [(8.74, 4.80), (8.52, 4.80)])
    arrow_line(s, [(4.68, 4.80), (4.46, 4.80)])
    text(s, "When CVs or requirements change, refresh the saved job and review the updated results.", .70, 6.05, 11.94, .50, size=18, color="violet")
    notes(s, "The numbered sequence runs across the upper row, then down and back across the lower row. An Express search can support the initial search without creating a vacancy record. Save it as a job when you need persistent reviews or a hiring record. Recruiter outreach and interviews take place through the organization's usual tools. After new CVs arrive, refresh the saved job; existing reviews and notes remain available.")

    standard(prs, 5, "Importing the CV library", "Supported inputs: digital PDF, DOCX, Markdown, and pasted text.", "02-library.png", [
        ("Open Library", "Choose Add files or Import folder."),
        ("Import the CVs", "Select documents or use Paste text."),
        ("Check processing", "Review import notices and candidate details."),
    ], "Files can also be dragged into the application. Arabic and English CVs can be searched with postings in either language. PDFs require a usable text layer; image-only scans need conversion outside Qumaa. Exact duplicate imports may be skipped. Inspect possible duplicates before linking records. The demonstration uses fictional CV text entered into an isolated library.", "Local matching requires no API key.")

    standard(prs, 6, "Creating a saved job", "Use a saved job for an assignment that requires ongoing records.", "03-job-editor.png", [
        ("Select New job", "Open Jobs and choose + New job."),
        ("Enter the posting", "Add the role, requirements, and qualifications."),
        ("Save the job", "The new job starts as Open with a timestamp."),
    ], "The editor separates Position, Job description, Requirements, and Education and certifications. Add criteria before saving or return through Edit job. Newly created jobs start Open and record their creation time. Use Change status after creation to put a job On hold, mark it Filled, or archive it. A saved job retains its posting, search results, reviews, shortlist, and status history on this device.", "Status changes use the Change status control.")

    s = slide_base(prs, "Express search", 7, "Paste a job and its requirements for a temporary search of the CV library.")
    steps(s, [
        ("Open Express search", "In Jobs, use the button beside + New job."),
        ("Paste the details", "Enter a description and requirements; find matches."),
        ("Review the results", "Discard the session or choose Save as job."),
    ])
    shot(s, "13-express-form.png", 4.05, 1.97, 8.66, 4.77, caption="Temporary search form")
    callout(s, "Nothing is saved unless you choose Save as job.")
    notes(s, "Express search is in Jobs next to New job. Paste the job description and requirements, then select Find matches. The search uses CVs already in the local library; it does not fetch applicants from external services. The temporary posting and results stay in memory. You can view candidate details or export the top three matches. Reviews, shortlists, and comparison become available after Save as job. Navigating between tabs can retain the current session, but it is not a saved vacancy. Discard or close the application to remove it. Save as job explicitly creates a persistent job. Optional AI reranking is available here under the same settings and data-review controls as a saved job.")

    standard(prs, 8, "Defining assessment criteria", "Separate essential requirements from preferred qualifications.", "04-criteria.png", [
        ("Create criteria", "Use requirement lines or select Add criterion."),
        ("Set importance", "Choose Required or Preferred for each item."),
        ("Use clear evidence", "Example: led training across restaurant sites."),
    ], "Criterion text contributes to the matching query. Required and Preferred organize recruiter assessment; they do not automatically exclude a candidate. The demonstration criteria cover multi-site training, onboarding design and measurement, and bilingual delivery. Save the job after editing. If criterion wording changes, revisit earlier assessments to ensure they still address the requirement.", "Required criteria do not automatically reject candidates.")

    standard(prs, 9, "Finding and filtering candidates", "Use the posting to search, then narrow the displayed results.", "05-filters.png", [
        ("Find matches", "Search from the selected job."),
        ("Inspect the ranking", "Use scores to prioritize evidence review."),
        ("Apply filters", "Filter by location, skills, and more."),
    ], "After the first search, Find matches becomes Refresh matches. Filters use extracted CV information plus recruiter profile corrections. Options include current or preferred location, experience, skills, tags, and availability. A text search can locate names, contact details, skills, and tags. Review-state filters help organize work. Filters change what is displayed; they do not verify candidate claims or change saved decisions.", "Profile corrections are used by screening filters.")

    side_panel(prs, 10, "Interpreting scores and evidence", "Inspect the underlying CV passages before making an assessment.", "06-evidence.png", [
        ("Open a candidate", "Select the candidate name in the results."),
        ("Inspect the evidence", "Read Score breakdown, Matching passages, and CV text."),
        ("Identify questions", "Record relevant gaps for screening or interview."),
    ], "Local scores combine semantic similarity, keyword overlap, and detected skills. They describe fit within the current search context and are not a probability of success. A high score still requires a review of scope, relevance, and supporting evidence. Missing words may reflect wording or omitted details. Scores can change when the posting or library changes.", "A match score is a screening aid, not a hiring decision.")

    standard(prs, 11, "Recording candidate reviews", "Job-specific reviews preserve assessments and follow-up questions.", "07-review.png", [
        ("Open Review", "Choose a review state and enter job notes."),
        ("Assess each criterion", "Record evidence, confirmation, or clarification."),
        ("Save the review", "Save changes before reviewing another person."),
    ], "Review states are Unreviewed, Maybe, Shortlisted, Contacted, and Not proceeding. Criterion assessments are Unreviewed, Evidence found, Needs clarification, Not stated, and Recruiter confirmed. Evidence found records supporting CV text; Recruiter confirmed records the recruiter's verification. Contacted is set manually after outreach through other tools. Keep decision reasons specific to the job.", "Job notes remain attached to this vacancy.")

    side_panel(prs, 12, "Maintaining candidate profiles", "Profiles retain screening information across saved jobs.", "08-profile.png", [
        ("Open the profile", "Choose Add details or Edit profile."),
        ("Record verified details", "Add corrections, salary, notice, availability, and tags."),
        ("Save the profile", "Use job reviews for vacancy-specific decisions."),
    ], "Recruiter corrections are stored separately from CV extraction. A blank correction uses the extracted value. The profile can hold expected salary, notice period, availability, preferred location, tags, and general notes. Profile changes apply across jobs. Job-specific recommendations and assessments belong in a job review. Verify information before relying on it for screening.", "Profile notes apply across jobs.")

    s = slide_base(prs, "Comparing candidates", 13, "Select two to four candidates and choose Compare candidates.")
    shot(s, "09-comparison-header.png", .66, 1.94, 12.04, 1.05, framed=False)
    shot(s, "09-comparison-criteria.png", .66, 3.24, 12.04, 2.96,
         caption="Criteria and supporting evidence from the same comparison", framed=False)
    text(s, "Review the same criteria, screening details, and unresolved questions for each person.", .75, 6.44, 11.82, .42, size=17, color="soft")
    notes(s, "Use Compare checkboxes in Candidates or Shortlist, then open Compare candidates. The comparison includes criterion assessments, suggested passages, recruiter profiles, and notes. Salary, notice, and availability appear after they are recorded. Not confirmed identifies information to verify. Use Read evidence for context. The two images are separate excerpts from the same actual comparison view.")

    standard(prs, 14, "Preparing and exporting a shortlist", "Choose candidates and set the order for the hiring review.", "10-shortlist.png", [
        ("Choose candidates", "Set relevant candidates to Shortlisted."),
        ("Arrange the shortlist", "Use the up and down arrows in Shortlist."),
        ("Export the record", "Select Export PDF or Export CSV."),
    ], "Exports contain the full ordered shortlist, including candidates hidden by a current display filter, together with job-specific recruiter notes. Contacted candidates remain on the shortlist. Refresh stale matches before exporting. PDF supports a review conversation; CSV can be opened in a spreadsheet. Exporting creates a file and does not send outreach.", "Exports include the full ordered shortlist.")

    side_panel(prs, 15, "Linking updated CV versions", "Confirm identity before linking two records as the same person.", "11-versions.png", [
        ("Import the new CV", "Add the updated document to Library."),
        ("Open Resume history", "Select Link another CV in the candidate panel."),
        ("Confirm the link", "Choose the record and select Link as same person."),
    ], "Linking connects recruiter context across CV versions. Check the identity before linking; similar experience alone is insufficient. Resolve conflicting recruiter details when prompted. Unlinking retains a copy of shared profile information and job reviews with each resulting candidate, allowing independent editing afterward. Search results may need refreshing after library changes.", "Linked versions retain the candidate's recruiting context.")

    standard(prs, 16, "Refreshing saved matches", "Refresh when the library or job requirements change.", "12-refresh.png", [
        ("Reopen the job", "Return after importing or updating CVs."),
        ("Read the notice", "Check whether the saved matches are out of date."),
        ("Refresh matches", "Review the updated ranking and evidence."),
    ], "New CVs, updated library information, and edited job requirements can make saved matches stale. Refresh recalculates local matches while retaining saved decisions and notes. Reassess criteria when their wording has changed. Restoring a backup also requires fresh searches after indexing finishes. If AI reranking was previously used, a local refresh creates a new local ranking that can be reranked again.", "Saved reviews and notes remain available after refresh.")

    standard(prs, 17, "Changing job status", "Each recorded status change includes a date and time.", "16-status.png", [
        ("Open Change status", "Use the control in the selected job."),
        ("Choose the status", "Open, On hold, Filled, or Archived."),
        ("Check the history", "Review the creation and status-change timestamps."),
    ], "New jobs begin Open with a creation event. A saved status change appends a timestamped entry to the job history. Selecting Filled also requires a hired candidate. Archiving or reopening a job retains its last recorded hiring outcome. Older jobs may have no historical events from before the feature existed; the application does not invent those dates. Use this record to understand when the vacancy changed state.", "Archived jobs retain their recorded hiring outcome.")

    standard(prs, 18, "Recording a hiring outcome", "A filled job identifies the hired person and optional backup candidates.", "17-hiring.png", [
        ("Select Filled", "Change the job status after confirming the outcome."),
        ("Choose the candidates", "Select the hired person and backups."),
        ("Save and inspect", "Review the recorded cards, timestamp, and linked CVs."),
    ], "Choose candidates already present in Library. A filled job requires a hired candidate. Up to two backup candidates can be recorded, in order: #2 and #3. Choose #2 before #3, and use a different person in each position. The saved cards display the recorded names and provide View CV links. Changes to hiring details receive their own history entry. Reopening or archiving preserves the last hiring outcome. Deleted CV records are shown as unavailable rather than exposing removed candidate details.", "Backup #3 requires backup #2.")

    s = slide_base(prs, "Selecting an AI provider and model", 19, "AI reranking is optional and configured in Settings.")
    for i, (provider, models) in enumerate([
        ("Anthropic", "Sonnet 5\nOpus 5\nFable 5.1"),
        ("OpenAI", "GPT-4o mini\nGPT-5.6\nGPT-6 Astra"),
    ]):
        y = 1.96 + i * 2.44
        rect(s, .66, y, 3.26, 2.20, "white", "line", True)
        text(s, provider, .87, y + .19, 2.80, .37, size=23, bold=True, color="violet")
        text(s, models, .87, y + .85, 2.80, 1.18, size=20)
    shot(s, "18-ai-settings.png", 4.18, 1.96, 8.53, 3.40)
    text(s, "Choose AI provider and AI model, add your own provider key, then enable reranking.", 4.34, 5.70, 8.12, .82, size=21, color="soft")
    notes(s, "Settings provides separate provider and model selectors. Anthropic choices are Sonnet 5, Opus 5, and Fable 5.1. OpenAI choices are GPT-4o mini, GPT-5.6, and GPT-6 Astra. Each provider has its own protected API key and remembered model selection. A key is not supplied with the installer. Switching to a provider without a key disables reranking until configured. Local matching remains available without either provider. This screenshot uses a fictional fixture; no live provider request was made.")

    side_panel(prs, 20, "Reviewing data before AI reranking", "The notice identifies the selected provider and the data to be sent.", "19-ai-review.png", [
        ("Start from results", "Choose Rerank with the selected model."),
        ("Read the summary", "Open the eye icon, See exported data, for request details."),
        ("Confirm the request", "Approve only after reviewing the information."),
    ], "The application first explains the data policy for the selected provider and model. The exact outbound JSON is collapsed under See exported data. Policy approval is remembered for that provider and model until the policy changes; each rerank still requires a separate native confirmation before sending. The reviewed request includes the posting and redacted matching passages from up to 15 candidates. Common contact details are masked, but redaction may not remove every identifying detail. Canceling the review leaves the local matches available.", "Each rerank requires confirmation before sending.")

    s = slide_base(prs, "Understanding an AI rerank", 21, "The selected model reorders a limited candidate set using the supplied evidence.", footer_label="AI reranking · Editable diagram")
    flow = [
        ("Local results", "The application prepares up to 15 candidates from the local search."),
        ("Reviewed request", "After confirmation, the selected provider receives the posting and redacted passages."),
        ("Returned assessment", "The model returns an order and a short evidence-based reason for each candidate."),
    ]
    for i, (heading, body) in enumerate(flow):
        x = .66 + i * 4.10
        rect(s, x, 1.98, 3.79, 2.50, "white", "line", True)
        text(s, heading, x + .22, 2.19, 3.35, .52, size=23, bold=True, color="violet")
        text(s, body, x + .22, 2.93, 3.34, 1.37, size=18)
    arrow_line(s, [(4.47, 3.12), (4.73, 3.12)])
    arrow_line(s, [(8.57, 3.12), (8.83, 3.12)])
    text(s, "Assessment basis", .79, 4.79, 5.64, .41, size=22, bold=True)
    text(s, "Relevant experience, demonstrated responsibility, and evidence for job requirements.", .79, 5.40, 5.51, 1.18, size=20, color="soft")
    text(s, "Reading the result", 6.89, 4.79, 5.54, .41, size=22, bold=True)
    text(s, "Candidate membership and local scores remain unchanged. The model label identifies who produced the order.", 6.89, 5.40, 5.50, 1.19, size=20, color="soft")
    notes(s, "All configured models use the same system instructions. They rank every supplied candidate, judge only the provided passages, preserve opaque candidate IDs, treat the payload as evidence rather than instructions, and return one reason of at most 25 words per candidate. The instructions emphasize demonstrated scope and relevance over confident wording. The request contains up to 2,500 characters of redacted passages per candidate. Reranking does not add candidates or replace local scores. Check the supporting CV evidence and verify missing information. An AI explanation is a screening aid; the recruiter remains responsible for the assessment.")

    standard(prs, 22, "Using a watched folder", "Import incoming CVs from a selected local folder while Qumaa is running.", "14-watched-folder.png", [
        ("Open Settings", "Find the Watched folder section."),
        ("Choose the folder", "Select the local folder used for incoming CVs."),
        ("Review imports", "Check Library and refresh relevant saved jobs."),
    ], "The watched folder monitors incoming supported documents while the application is running. Keep original documents organized separately. Import notices and duplicate handling still apply. New imports do not automatically update every saved ranking; refresh the relevant job. Watched-folder configuration is specific to this device and is not included in a portable backup.", "Refresh saved matches after new imports.")

    standard(prs, 23, "Backing up and restoring records", "Use Settings to create or restore a portable library backup.", "14-backups.png", [
        ("Create a backup", "In Library & backups, select Back up."),
        ("Store it securely", "Backups contain candidate records."),
        ("Restore when needed", "After indexing finishes, refresh saved matches."),
    ], "Portable backups include resume text, saved jobs, profiles, reviews, hiring records, and portable preferences. They exclude API keys, consent approvals, watched-folder configuration, original-file paths, and window position. They are not encrypted. Keep backups in an appropriate secure location and retain original CV files separately if needed. Restoring replaces the current local library after confirmation and rebuilds the search index. Reconfigure device-specific settings afterward.", "Backups exclude API keys and original CV files.")

    s = slide_base(prs, "Reference: where information is kept", 24, "Use the appropriate record for the information you need to retain.", footer_label="Recruiter guide")
    rows = [
        ("Library and candidate profile", "CV text, linked versions, corrected details, and general screening notes."),
        ("Saved job", "Posting, criteria, results, candidate reviews, shortlist, status history, and hiring outcome."),
        ("Express search", "Temporary posting and search work; retained only if you choose Save as job."),
        ("Optional provider request", "Reviewed posting and redacted passages, sent only after confirmation."),
    ]
    for i, (heading, body) in enumerate(rows):
        y = 1.97 + i * 1.02
        rect(s, .67, y, 12.0, .91, "white", "line", True)
        text(s, heading, .87, y + .19, 3.74, .58, size=18, bold=True, color="violet")
        text(s, body, 4.95, y + .18, 7.46, .61, size=18, color="soft")
    text(s, "Prepared by Fawaz Aljathnan", .83, 6.33, 7.15, .32, size=16, color="muted", link=LINKEDIN)
    text(s, "LinkedIn profile", 9.35, 6.34, 3.10, .30, size=15, color="violet", align=PP_ALIGN.RIGHT, link=LINKEDIN)
    notes(s, "This table distinguishes information attached to a candidate, information attached to a saved vacancy, a temporary search session, and an optional external request. Use profile notes for information relevant across jobs and review notes for vacancy-specific decisions. All application images in this guide use fictional demonstration records. Prepared by Fawaz Aljathnan. LinkedIn: https://www.linkedin.com/in/fawazsalehq/.")

    names = ["introduction", "contents", "workspace", "workflow", "library", "saved-job", "express-search", "criteria", "matches", "evidence", "reviews", "profiles", "comparison", "shortlist", "versions", "refresh", "status", "hiring", "ai-settings", "ai-review", "ai-workflow", "watched-folder", "backups", "reference"]
    assert len(prs.slides) == len(names) == 24
    for slide, name in zip(prs.slides, names):
        slide._element.cSld.set("name", name)
    for shape, page_number in toc_links:
        shape.click_action.target_slide = prs.slides[page_number - 1]
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return prs


def write_reports(prs):
    """Keep the editable deck, storyboard, and build inventory in agreement."""
    import hashlib
    import json

    lines = ["# Qumaa recruiter guide: English storyboard", "",
             "24 pages. Current application screenshots use fictional recruiting data.",
             "The contents page follows the introduction. Copy and notes are extracted from the editable presentation.", ""]
    pages = []
    for number, slide in enumerate(prs.slides, 1):
        copy = [shape.text for shape in slide.shapes if shape.has_text_frame
                and shape.top < Inches(7) and shape.text.strip() and shape.text != "QUMAA"]
        pictures = [shape.name.split(":", 1)[1] for shape in slide.shapes
                    if shape.name.startswith("screenshot:")]
        lines += [f"## {number:02d}. {copy[0]}", "", "Page copy:", ""]
        lines += ["- " + value.replace("\n", " / ") for value in copy[1:]]
        visuals = ", ".join("`" + name + "`" for name in pictures) if pictures else "editable presentation shapes and text"
        lines += ["", "Visuals: " + visuals + ".", "", "Speaker notes:", "",
                  slide.notes_slide.notes_text_frame.text, ""]
        pages.append({"page": number, "name": slide._element.cSld.get("name"),
                      "title": copy[0], "screenshots": pictures, "speakerNotes": True})
    (BASE / "source" / "storyboard.md").write_text("\n".join(lines))
    (BASE / "source" / "build-report.json").write_text(json.dumps({
        "slides": len(prs.slides), "language": "en", "contentsPage": 2,
        "sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(), "pages": pages,
    }, indent=2) + "\n")


def main():
    global ALLOW_MISSING
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--allow-missing", action="store_true", help="Use labelled placeholders for missing images (development only).")
    args = parser.parse_args()
    ALLOW_MISSING = args.allow_missing
    missing = [name for name in REQUIRED if not (SHOTS / name).is_file()]
    if missing and not ALLOW_MISSING:
        parser.error("Missing actual application screenshots: " + ", ".join(missing))
    prs = make_deck()
    if not ALLOW_MISSING:
        write_reports(prs)
    print(f"Created {OUTPUT} ({len(prs.slides)} slides)")
    if missing:
        print("DEVELOPMENT OUTPUT: contains screenshot placeholders; do not distribute.")


if __name__ == "__main__":
    main()
