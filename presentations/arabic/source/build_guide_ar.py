"""Build the Arabic guide from the shared slide templates and actual Arabic captures.

All visible text and speaker notes must have an explicit translation. Diagram
mirroring uses semantic slide names, so inserting a contents slide or expanding
the guide cannot accidentally rearrange an unrelated screen.
"""
from __future__ import annotations

import hashlib
import importlib.util
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

BASE = Path(__file__).resolve().parents[1]
ENGLISH = BASE.parent / "source" / "build_guide.py"
sys.dont_write_bytecode = True
spec = importlib.util.spec_from_file_location("english_guide", ENGLISH)
guide = importlib.util.module_from_spec(spec)
spec.loader.exec_module(guide)
TRANSLATIONS = json.loads((BASE / "source" / "translations.json").read_text())
USED = set()
ARABIC = re.compile(r"[\u0600-\u06ff]")
NUMERALS = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")


def translate(value):
    if value in TRANSLATIONS:
        USED.add(value)
        result = TRANSLATIONS[value]
    elif value.isdigit():
        result = value.translate(NUMERALS)
    elif re.fullmatch(r"Qumaa (v\S+)\s+/\s+(\d+)", value):
        version, number = re.fullmatch(r"Qumaa (v\S+)\s+/\s+(\d+)", value).groups()
        result = f"قُمع {version}   /   {number.translate(NUMERALS)}"
    elif re.fullmatch(r"\d+[–-]\d+", value):
        result = value.translate(NUMERALS)
    elif value in {"Fawaz Aljathnan", "in"}:
        result = value
    else:
        raise ValueError(f"Missing Arabic translation: {value!r}")
    if "\u2014" in result:
        raise ValueError("An em dash is not permitted in the Arabic presentation.")
    return result


def set_language(shape, rtl):
    for paragraph in shape.text_frame.paragraphs:
        paragraph._p.get_or_add_pPr().set("rtl", "1" if rtl else "0")
        for run in paragraph.runs:
            props = run._r.get_or_add_rPr()
            props.set("lang", "ar-SA" if rtl else "en-US")
            if rtl:
                props.set("altLang", "en-US")
                # Explicit complex-script font avoids a viewer-dependent Arabic
                # fallback. Reapplying language should not duplicate XML nodes.
                existing = props.find("{http://schemas.openxmlformats.org/drawingml/2006/main}cs")
                if existing is None:
                    existing = OxmlElement("a:cs")
                    props.append(existing)
                existing.set("typeface", "Arial")


original_text = guide.text
original_notes = guide.notes
original_number_badge = guide.number_badge


def text(slide, value, *args, **kwargs):
    rendered = translate(value)
    rtl = bool(ARABIC.search(rendered))
    # Explicit direction also protects punctuation in Office preview importers
    # that ignore DrawingML's rtl flag when converting Arabic runs to PDF.
    if rtl:
        rendered = "\n".join("\u202b" + line + "\u202c" for line in rendered.split("\n"))
    # Preserve the original frame positions. Arabic paragraphs read from their
    # right edge; centered badges and the compact brand/footer keep alignment.
    if rtl and kwargs.get("align", PP_ALIGN.LEFT) != PP_ALIGN.CENTER:
        kwargs["align"] = PP_ALIGN.RIGHT
    if value in {"QUMAA", "Actual application · Fictional demo data", "Recruiter workflow · Editable diagram", "Recruiter guide"}:
        kwargs["align"] = PP_ALIGN.LEFT
    shape = original_text(slide, rendered, *args, **kwargs)
    set_language(shape, rtl)
    return shape


def notes(slide, value):
    rendered = translate(value)
    rendered = "\n".join("\u202b" + line + "\u202c" for line in rendered.split("\n"))
    original_notes(slide, rendered)
    for paragraph in slide.notes_slide.notes_text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.RIGHT
        paragraph._p.get_or_add_pPr().set("rtl", "1")
        for run in paragraph.runs:
            run.font.name = "Arial"
            run._r.get_or_add_rPr().set("lang", "ar-SA")


def number_badge(slide, value, *args, **kwargs):
    numeral = translate(value)
    badge = original_number_badge(slide, numeral, *args, **kwargs)
    set_language(badge, bool(ARABIC.search(numeral)))
    return badge


def steps(slide, items, x=.63, y=1.93, w=3.08, gap=1.35, size=19):
    # Mirror only the badge and text within the existing instruction column.
    for i, (heading, body) in enumerate(items):
        yy = y + i * gap
        badge_x = x + w - .34
        numeral = translate(str(i + 1))
        badge = guide.number_badge(slide, numeral, badge_x, yy + .025)
        set_language(badge, bool(ARABIC.search(numeral)))
        text(slide, heading, x, yy, w - .47, .64, size=size, bold=True)
        text(slide, body, x, yy + .64, w - .47, gap - .60, size=15, color="soft")


def mirror_body(slide, slide_width):
    """Mirror native diagram/card positions while keeping text and images intact."""
    for shape in slide.shapes:
        if Inches(1.90) <= shape.top < Inches(6.70):
            shape.left = slide_width - shape.left - shape.width
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                transform = shape._element.spPr.xfrm
                transform.set("flipH", "0" if transform.get("flipH") == "1" else "1")


def plain_text(value):
    return re.sub(r"[\u202a-\u202e\u2066-\u2069]", "", value).strip()


def write_storyboard(prs):
    """Keep the readable Arabic outline synchronized with delivered slide copy."""
    lines = [
        "# قُمع: دليل مسؤول التوظيف",
        "",
        f"دليل عملي من {str(len(prs.slides)).translate(NUMERALS)} شريحة باللغة العربية، بنسبة عرض ١٦:٩.",
        "",
        "تظهر المحتويات في الشريحة الثانية. النصوص والمخططات قابلة للتعديل، وتستخدم اللقطات واجهة التطبيق العربية وبيانات توضيحية خيالية.",
        "",
        "يُولَّد هذا الملف من النصوص وملاحظات المقدّم في العرض النهائي لضمان تطابق المحتوى.",
        "",
    ]
    for index, slide in enumerate(prs.slides, 1):
        text_shapes = [shape for shape in slide.shapes if shape.has_text_frame and plain_text(shape.text)]
        title = max(text_shapes, key=lambda shape: max(
            (run.font.size or 0 for paragraph in shape.text_frame.paragraphs for run in paragraph.runs), default=0
        ))
        lines += [f"## {str(index).zfill(2).translate(NUMERALS)}. {plain_text(title.text).replace(chr(10), ' ')}", "", "**المحتوى الظاهر**", ""]
        for shape in text_shapes:
            if shape is title or shape.top < Inches(1.20) or shape.top >= Inches(7.0):
                continue
            lines.append("- " + plain_text(shape.text).replace("\n", " / "))
        lines += ["", "**ملاحظات المقدّم**", "", plain_text(slide.notes_slide.notes_text_frame.text), ""]
    (BASE / "source" / "storyboard-ar.md").write_text("\n".join(lines), encoding="utf-8")


def main():
    guide.BASE = BASE
    guide.SHOTS = BASE / "screenshots"
    guide.OUTPUT = BASE / "Qumaa_Recruiter_Guide_AR.pptx"
    guide.text, guide.notes, guide.steps, guide.number_badge = text, notes, steps, number_badge
    # Each capture is already framed around the relevant Arabic UI region. The
    # deck embeds the real PNG without language-specific, stale pixel crops.
    capture = json.loads((guide.SHOTS / "manifest.json").read_text())
    assert capture["language"] == "ar" and capture["direction"] == "rtl"
    missing = [name for name in guide.REQUIRED if not (guide.SHOTS / name).is_file()]
    if missing:
        raise FileNotFoundError(f"Missing real Arabic screenshots: {missing}")
    prs = guide.make_deck()
    prs.core_properties.title = "قُمع | دليل مسؤول التوظيف"
    prs.core_properties.subject = "دليل عملي للبحث عن المرشحين ومراجعتهم باستخدام قُمع"
    prs.core_properties.keywords = "قُمع، التوظيف، السير الذاتية، الوظائف، القائمة المختصرة، العربية"
    prs.core_properties.comments = "لقطات من واجهة التطبيق العربية باستخدام بيانات توضيحية خيالية."
    slide_names = [slide._element.cSld.get("name") for slide in prs.slides]
    if len(slide_names) != len(set(slide_names)) or None in slide_names:
        raise ValueError("Every slide must have a unique semantic name in the shared template.")
    for slide, name in zip(prs.slides, slide_names):
        if name in {"workflow", "contents", "ai-workflow", "reference"}:
            mirror_body(slide, prs.slide_width)
    if slide_names[1] != "contents":
        raise ValueError("The table of contents must be slide 2.")
    prs.save(guide.OUTPUT)
    with ZipFile(guide.OUTPUT) as package:
        for name in package.namelist():
            if name.endswith(".xml") and "\u2014" in package.read(name).decode("utf-8"):
                raise ValueError(f"Em dash found in presentation XML: {name}")
    write_storyboard(prs)
    report = {
        "slides": len(prs.slides), "language": "ar-SA", "emDashCount": 0,
        "translatedStrings": len(USED), "sourceTemplate": str(ENGLISH),
        "slideNames": slide_names, "contentsSlide": 2,
        "editablePipeline": True, "actualArabicScreenshots": True,
        "sha256": hashlib.sha256(guide.OUTPUT.read_bytes()).hexdigest(),
    }
    (BASE / "source" / "build-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    print(guide.OUTPUT)


if __name__ == "__main__":
    main()
